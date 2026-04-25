"""Validate the solver against an analytically-tractable hull.

A wall-sided rectangular block of length L, beam B, depth D has the
following exact properties at draft T (T ≤ D):

    ∇  = L · B · T
    Aw = L · B
    LCB = LCF = L / 2          (with x measured from one end)
    KB = T / 2
    IT  = L · B³ / 12          (transverse MoI of the waterplane about CL)
    BMt = IT / ∇ = B² / (12 · T)
    KMt = T/2 + B² / (12·T)

For a box hull at moderate heel φ (deck-edge not yet immersed and bilge
not emergent), the wall-sided approximation gives:

    GZ(φ) ≈ (GMt + ½ · BMt · tan²φ) · sinφ
    KN(φ) = GZ(φ) + KG · sin φ

We use a synthetic offsets table (constant half-breadth = B/2) and
check that the solver matches these numbers to a tight tolerance.
Also writes a sample input workbook ``data/sample_box.xlsx`` that can be
re-used as a template for the real input file.
"""

from __future__ import annotations
import sys
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from hydrohack.io_offsets import load_explicit, load_workbook  # noqa: E402
from hydrohack.hydrostatics import compute_upright              # noqa: E402
from hydrohack.stability import gz_curve                        # noqa: E402


def make_box_input(L=100.0, B=12.0, D=8.0, T=4.0, rho=1.025, KG=3.0,
                   n_stations=21, n_waterlines=9):
    """Return ShipInput for a rectangular wall-sided block hull."""
    offsets = np.full((n_stations, n_waterlines), B / 2.0)
    return load_explicit(
        offsets=offsets,
        station_spacing=L / (n_stations - 1),
        waterline_spacing=D / (n_waterlines - 1),
        LBP=L, B=B, D=D, T=T, rho=rho, KG=KG,
    )


def write_sample_xlsx(path: Path, L=100.0, B=12.0, D=8.0, T=4.0, rho=1.025, KG=3.0,
                      n_stations=21, n_waterlines=9):
    """Write a sample input workbook with a Particulars sheet + Offsets sheet."""
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()

    # --- Particulars sheet ---
    ws = wb.active
    ws.title = "Particulars"
    rows = [
        ("Parameter", "Value", "Unit"),
        ("LBP", L, "m"),
        ("Beam", B, "m"),
        ("Depth", D, "m"),
        ("Draft", T, "m"),
        ("Density", rho, "t/m^3"),
        ("KG", KG, "m"),
        ("Station spacing", L / (n_stations - 1), "m"),
        ("Waterline spacing", D / (n_waterlines - 1), "m"),
        ("n_stations", n_stations, ""),
        ("n_waterlines", n_waterlines, ""),
    ]
    for r in rows:
        ws.append(r)

    # --- Offsets sheet: rows = stations, cols = waterlines, header rows/cols ---
    ws2 = wb.create_sheet("Offsets")
    z = np.linspace(0, D, n_waterlines)
    x = np.linspace(0, L, n_stations)
    ws2.append(["x \\ z"] + [float(zi) for zi in z])
    for i, xi in enumerate(x):
        ws2.append([float(xi)] + [B / 2.0] * n_waterlines)

    wb.save(path)


def expected_box(L, B, D, T, rho, KG):
    V = L * B * T
    return {
        "V":   V,
        "Aw":  L * B,
        "LCB": L / 2,
        "LCF": L / 2,
        "KB":  T / 2,
        "IT":  L * B**3 / 12.0,
        "BMt": B**2 / (12.0 * T),
        "KMt": T / 2.0 + B**2 / (12.0 * T),
        "GMt": T / 2.0 + B**2 / (12.0 * T) - KG,
    }


def run_validation():
    L, B, D, T, rho, KG = 100.0, 12.0, 8.0, 4.0, 1.025, 3.0
    ship = make_box_input(L, B, D, T, rho, KG)
    up = compute_upright(ship.hull, ship.T, ship.rho, ship.KG)
    exp = expected_box(L, B, D, T, rho, KG)

    print("Box-hull validation  (L=100, B=12, D=8, T=4, ρ=1.025, KG=3)")
    print("-" * 72)
    rows = [
        ("V (∇)",       up.V,    exp["V"]),
        ("Aw",          up.Aw,   exp["Aw"]),
        ("LCB",         up.LCB,  exp["LCB"]),
        ("LCF",         up.LCF,  exp["LCF"]),
        ("KB",          up.KB,   exp["KB"]),
        ("IT",          up.IT,   exp["IT"]),
        ("BMt",         up.BMt,  exp["BMt"]),
        ("KMt",         up.KMt,  exp["KMt"]),
        ("GMt",         up.GMt,  exp["GMt"]),
    ]
    print(f"{'Quantity':>10}  {'computed':>14}  {'expected':>14}  {'rel.err':>10}")
    max_err = 0.0
    for name, c, e in rows:
        err = abs(c - e) / max(abs(e), 1e-9)
        max_err = max(max_err, err)
        print(f"{name:>10}  {c:>14.6f}  {e:>14.6f}  {err:>10.2e}")
    print("-" * 72)
    print(f"max relative error: {max_err:.2e}")

    # Heeled check at small angle: wall-sided formula
    angles = np.array([0.0, 2.0, 5.0, 10.0, 15.0, 20.0])
    curve = gz_curve(ship.hull, ship.T, ship.rho, ship.KG, angles_deg=angles)
    print("\nGZ small-angle check vs wall-sided formula:")
    GMt, BMt = exp["GMt"], exp["BMt"]
    for phi, gz_calc in zip(angles, curve.GZ):
        phi_r = np.deg2rad(phi)
        gz_exp = (GMt + 0.5 * BMt * np.tan(phi_r) ** 2) * np.sin(phi_r)
        print(f"  φ={phi:5.1f}°   GZ_calc={gz_calc:+8.5f}   GZ_wall-sided={gz_exp:+8.5f}   "
              f"Δ={gz_calc - gz_exp:+8.2e}")

    # Also export a sample workbook and try round-trip via load_workbook
    sample_path = ROOT / "data" / "sample_box.xlsx"
    write_sample_xlsx(sample_path, L, B, D, T, rho, KG)
    ship2 = load_workbook(sample_path)
    up2 = compute_upright(ship2.hull, ship2.T, ship2.rho, ship2.KG)
    print(f"\nRound-trip via xlsx ({sample_path}):  ∇={up2.V:.4f}  KMt={up2.KMt:.4f}  GMt={up2.GMt:.4f}")

    assert max_err < 1e-3, f"Upright validation failed: max_err={max_err:.2e}"
    print("\n✓ Validation passed.")


if __name__ == "__main__":
    run_validation()
